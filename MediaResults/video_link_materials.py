import os
import io
import pandas as pd
import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils.cell import get_column_letter

class MediaVideoLink:
    def __init__(self, t_engine, oc_client, year, raw_by_channel):
        self.t_engine   = t_engine
        self.oc_client  = oc_client
        self.year       = year
        self.file_name = f"Video Materials Media {year}.xlsx" 
        self.raw = raw_by_channel
    
    def decorate_table(self, sheet_name):
        wb = load_workbook(self.file_name)
        ws = wb[sheet_name]

        header_fill  = PatternFill("solid", fgColor="D3D3D3")
        thin_border  = Border(
            left=Side(style='thin'),  right=Side(style='thin'),
            top =Side(style='thin'),  bottom=Side(style='thin')
        )

        # Обрабатываем ссылки
        header = [c.value for c in ws[1]]
        if "video_link" in header:
            idx        = header.index("video_link") + 1
            col_letter = get_column_letter(idx)
            for row in range(2, ws.max_row + 1):
                cell = ws[f"{col_letter}{row}"]
                text = str(cell.value).strip() if cell.value else ""
                # если текст начинается с http/https — делаем hyperlink
                if text.lower().startswith(("http://", "https://")):
                    # если уже есть hyperlink, оставляем его
                    if not cell.hyperlink:
                        cell.hyperlink = text
                    cell.style = "Hyperlink"

        # 1) заголовки: серый фон + bold + граница
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.border = thin_border

        # 2) остальная сетка + авто-ширина
        for col in ws.columns:
            max_len = 0
            for cell in col:
                cell.border = thin_border
                max_len = max(max_len, len(str(cell.value or "")))
            ws.column_dimensions[col[0].column_letter].width = max_len + 2

        wb.save(self.file_name)

    def save_sheet(self, df: pd.DataFrame, sheet_name: str):
        file_exists = os.path.exists(self.file_name)
        mode = "a" if file_exists else "w"

        with pd.ExcelWriter(self.file_name,
                            engine="openpyxl",
                            mode=mode,
                            if_sheet_exists="replace" if file_exists else None) as writer:
            if file_exists:
                writer.book   = load_workbook(self.file_name)
                writer.sheets = {ws.title: ws for ws in writer.book.worksheets}
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            writer.save()

    def generate_raw_columns(self, channel):
        info_df = self.raw[channel]
        df = pd.DataFrame({
            "flight_name":                 info_df["flight_name"],
            "vertical":                    info_df["vertical"],
            "channel":                     channel,
            "date_start":                  info_df["date_start"],
            "date_end":                    info_df["date_end"],
            "video_link":                  None,
            "in_visualizer":               True,
        }).sort_values(["date_start", "date_end", "flight_name"])
        
        return df

    def make_final_table(self, df, channel: str, archive=False):
        """
        :channel: 'TV' | 'OOH' | 'Digital'
        """
        df = df.sort_values(["date_start", "date_end", "flight_name"]) \
               .reset_index(drop=True)

        # Приводим даты к типу date 
        for col in ("date_start", "date_end"):
            df[col] = pd.to_datetime(df[col]).dt.date

        if archive:
            sheet_name = f"Media Materials {channel} Archive"
        else:
            sheet_name = f"Media Materials {channel}"
            
            df = df.drop(columns="in_visualizer")

        self.save_sheet(df, sheet_name)
        self.decorate_table(sheet_name)

    def generate_excel(self):
        channels = ["TV", "OOH", "Digital"]
        
        for ch in channels:
            df = self.generate_raw_columns(ch)
            self.make_final_table(df, ch, archive=False)
            
        for ch in channels:
            raw = self.generate_raw_columns(ch)
            empty = raw.iloc[0:0].copy()
            self.make_final_table(empty, ch, archive=True)
            
        self.upload_to_nextcloud()

    def get_remote_parh(self, folder="MediaResults"):
        remote_path = f"{folder}/{self.file_name}"
        return remote_path

    def upload_to_nextcloud(self, folder="MediaResults"):
        remote_path = f"{folder}/{self.file_name}"
        self.oc_client.list(folder)
        self.oc_client.put_file(remote_path, self.file_name, chunked=False)

    def get_media_materials_table(self, sheet_name):
        remote_path = self.get_remote_parh()
        file_bytes = self.oc_client.get_file_contents(remote_path)
        with io.BytesIO(file_bytes) as buffer:
            df = pd.read_excel(buffer, sheet_name=sheet_name)
        return df

    def update_table(self, channel: str, folder="MediaResults"):
        # Сначала скачиваем текущий файл из Nextcloud, чтобы не потерять другие листы
        try:
            file_bytes = self.oc_client.get_file_contents(self.get_remote_parh(folder))
            with open(self.file_name, "wb") as f:
                f.write(file_bytes)
        except Exception:
            # файл может ещё не существовать на сервере
            pass
        
        main_sheet    = f"Media Materials {channel}"
        archive_sheet = f"Media Materials {channel} Archive"

        # Подтягиваем текущие DataFrame из NC
        main_df    = self.get_media_materials_table(main_sheet)
        main_df["in_visualizer"] = True
        archive_df = self.get_media_materials_table(archive_sheet)

        # Получаем свежие флайты из визуализатора
        vis_df = self.generate_raw_columns(channel).copy()

        # Обновляем основной лист
        active_set  = set(vis_df["flight_name"])
        current_set = set(main_df["flight_name"])

        # Строки, которых больше нет → в архив
        to_archive_mask = ~main_df["flight_name"].isin(active_set)
        if to_archive_mask.any():
            gone_rows = main_df[to_archive_mask].copy()
            gone_rows["in_visualizer"] = False
            archive_df = pd.concat([archive_df, gone_rows], ignore_index=True)
            main_df = main_df[~to_archive_mask]

        # Новые флайты → в основной
        new_names = active_set - current_set
        if new_names:
            new_rows = vis_df[vis_df["flight_name"].isin(new_names)].copy()
            main_df  = pd.concat([main_df, new_rows], ignore_index=True)

        self.make_final_table(main_df, channel)
        self.make_final_table(archive_df, channel, archive=True)
        
        self.upload_to_nextcloud(folder)
        
        return main_df